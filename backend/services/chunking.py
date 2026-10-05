def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    """Naive word-based chunking with overlap. Good enough for v1 — swap for a
    sentence-aware splitter later if retrieval quality needs it."""
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunks.append(" ".join(words[start:end]))
        start = end - overlap
    return chunks


def extract_pdf_text(file_bytes: bytes) -> str:
    from pypdf import PdfReader
    from io import BytesIO

    reader = PdfReader(BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_pptx_text(file_bytes: bytes) -> str:
    from pptx import Presentation
    from io import BytesIO

    def extract_from_shapes(shapes):
        parts = []
        for shape in shapes:
            if shape.shape_type == 6:  # GROUP shape — recurse into it
                parts.extend(extract_from_shapes(shape.shapes))
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        parts.append(text)
        return parts

    prs = Presentation(BytesIO(file_bytes))
    text_parts = []
    for slide in prs.slides:
        text_parts.extend(extract_from_shapes(slide.shapes))
    return "\n".join(text_parts)
